"""Renders reviewer / mock exam data into a Word file (Times New Roman 12), then to PDF via Word."""
import os
import random
import subprocess

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt

LETTERS = "ABCDEFGH"


def Q(q, correct, wrong, why=""):
    return {"kind": "q", "q": q, "correct": correct, "wrong": wrong, "why": why}


def M(title, pairs, extra=(), why=""):
    """Dropdown-style matching: pairs = [(term, definition)], extra = distractor definitions."""
    return {"kind": "m", "title": title, "pairs": pairs, "extra": list(extra), "why": why}


def L(title, answers, why="", ordered=False):
    """Typed-answer (enumeration) item: the student writes out len(answers) items."""
    return {"kind": "l", "title": title, "answers": answers, "why": why, "ordered": ordered}


def _style(doc):
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    for name in ("Heading 1", "Heading 2", "List Number", "Title"):
        s = doc.styles[name]
        s.font.name = "Times New Roman"
        rf = s.element.rPr.rFonts
        for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
            rf.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + attr, "Times New Roman")
        for attr in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
            rf.attrib.pop("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}" + attr, None)
        s.font.color.rgb = None
        if name.startswith("Heading"):
            s.font.bold = True
    doc.styles["Title"].font.size = Pt(20)
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 2"].font.size = Pt(12)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = sec.bottom_margin = Inches(1)


def build(path_docx, title, subtitle, sections, seed=7, answers_inline=False):
    rnd = random.Random(seed)
    doc = Document()
    _style(doc)
    doc.add_paragraph(title, style="Title")
    p = doc.add_paragraph(subtitle)
    p.runs[0].italic = True
    doc.add_paragraph(
        "Instructions: choose the best answer for every item. Matching items use a dropdown on the real "
        "exam, so pick one letter for each term. The answer key is at the end."
    )
    key = []  # (number, answer text, why)
    n = 0
    for heading, items in sections:
        doc.add_heading(heading, level=1)
        for it in items:
            if it["kind"] == "q":
                n += 1
                opts = [it["correct"]] + list(it["wrong"])
                rnd.shuffle(opts)
                ans = LETTERS[opts.index(it["correct"])]
                para = doc.add_paragraph(it["q"], style="List Number")
                para.paragraph_format.keep_with_next = True
                for i, o in enumerate(opts):
                    cp = doc.add_paragraph(f"{LETTERS[i]}.  {o}")
                    cp.paragraph_format.left_indent = Inches(0.6)
                    cp.paragraph_format.space_after = Pt(0)
                    if i < len(opts) - 1:
                        cp.paragraph_format.keep_with_next = True
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                key.append((n, f"{ans}. {it['correct']}", it["why"]))
            elif it["kind"] == "l":
                n += 1
                para = doc.add_paragraph(f"{it['title']} ({len(it['answers'])} items)", style="List Number")
                para.paragraph_format.keep_with_next = True
                for i in range(len(it["answers"])):
                    lp = doc.add_paragraph(f"{i + 1}.  ____________________________")
                    lp.paragraph_format.left_indent = Inches(0.6)
                    lp.paragraph_format.space_after = Pt(0)
                    lp.paragraph_format.keep_with_next = i < len(it["answers"]) - 1
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                sep = "; " if it["ordered"] else ", "
                tag = " (in order)" if it["ordered"] else " (any order)"
                key.append((n, sep.join(it["answers"]) + tag, it["why"]))
            else:
                n += 1
                terms = [t for t, _ in it["pairs"]]
                defs = [d for _, d in it["pairs"]] + it["extra"]
                order = list(range(len(defs)))
                rnd.shuffle(order)
                shown = [defs[i] for i in order]
                para = doc.add_paragraph(it["title"], style="List Number")
                para.paragraph_format.keep_with_next = True
                for i, d in enumerate(shown):
                    dp = doc.add_paragraph(f"{LETTERS[i]}.  {d}")
                    dp.paragraph_format.left_indent = Inches(0.6)
                    dp.paragraph_format.space_after = Pt(0)
                    dp.paragraph_format.keep_with_next = True
                for j, t in enumerate(terms):
                    tp = doc.add_paragraph(f"{t}:  ______")
                    tp.paragraph_format.left_indent = Inches(0.6)
                    tp.paragraph_format.space_after = Pt(0)
                doc.add_paragraph().paragraph_format.space_after = Pt(2)
                res = "; ".join(f"{t} = {LETTERS[shown.index(d)]}" for t, d in it["pairs"])
                key.append((n, res, it["why"]))
    doc.add_page_break()
    doc.add_heading("Answer Key", level=1)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    for c, h in zip(tbl.rows[0].cells, ("No.", "Answer", "Why")):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for num, a, w in key:
        r = tbl.add_row().cells
        r[0].text, r[1].text, r[2].text = str(num), a, w
    for row in tbl.rows:
        row.cells[0].width = Inches(0.5)
        row.cells[1].width = Inches(2.6)
        row.cells[2].width = Inches(3.4)
    doc.save(path_docx)
    return n, key


def to_pdf(path_docx):
    pdf = os.path.splitext(path_docx)[0] + ".pdf"
    ps = (
        "$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
        f"$d=$w.Documents.Open('{path_docx}');$d.SaveAs([ref]'{pdf}',[ref]17);$d.Close();$w.Quit()"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
    return pdf
