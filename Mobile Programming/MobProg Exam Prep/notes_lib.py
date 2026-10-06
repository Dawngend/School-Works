"""Renders study notes (readable, not quiz items) into a Word file, Times New Roman 12, then PDF via Word."""
import os
import subprocess

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _fonts(style):
    style.font.name = "Times New Roman"
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.append(rf)
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(W + a, "Times New Roman")
    for a in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"):
        rf.attrib.pop(W + a, None)


def _setup(doc):
    _fonts(doc.styles["Normal"])
    doc.styles["Normal"].font.size = Pt(12)
    for name, size in (("Title", 20), ("Heading 1", 15), ("Heading 2", 13), ("Heading 3", 12)):
        s = doc.styles[name]
        _fonts(s)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = None
    for name in ("List Bullet", "List Number", "List Bullet 2"):
        _fonts(doc.styles[name])
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Inches(1)
        sec.top_margin = sec.bottom_margin = Inches(1)


def _shade(cell, fill):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcpr.append(shd)


def _runs(par, text):
    """**bold** markup inside a string."""
    parts = text.replace("`", "").split("**")
    for i, part in enumerate(parts):
        if not part:
            continue
        r = par.add_run(part)
        r.bold = i % 2 == 1


class Notes:
    def __init__(self, title, subtitle):
        self.doc = Document()
        _setup(self.doc)
        self.doc.add_paragraph(title, style="Title")
        p = self.doc.add_paragraph(subtitle)
        p.runs[0].italic = True

    def h1(self, t):
        self.doc.add_heading(t, level=1)

    def h2(self, t):
        self.doc.add_heading(t, level=2)

    def h3(self, t):
        self.doc.add_heading(t, level=3)

    def p(self, t):
        par = self.doc.add_paragraph()
        _runs(par, t)
        par.paragraph_format.space_after = Pt(6)
        return par

    def bullets(self, items):
        for it in items:
            if isinstance(it, (list, tuple)):
                for sub in it:
                    par = self.doc.add_paragraph(style="List Bullet 2")
                    _runs(par, sub)
            else:
                par = self.doc.add_paragraph(style="List Bullet")
                _runs(par, it)

    def numbered(self, items):
        """Real numbered list that restarts at 1 each call."""
        num_id = self._new_num()
        for it in items:
            par = self.doc.add_paragraph(style="List Number")
            _runs(par, it)
            ppr = par._p.get_or_add_pPr()
            numpr = OxmlElement("w:numPr")
            il = OxmlElement("w:ilvl")
            il.set(qn("w:val"), "0")
            ni = OxmlElement("w:numId")
            ni.set(qn("w:val"), str(num_id))
            numpr.append(il)
            numpr.append(ni)
            ppr.append(numpr)

    def _new_num(self):
        numbering = self.doc.part.numbering_part.numbering_definitions._numbering
        base = None
        for st in self.doc.styles:
            if st.name == "List Number":
                ppr = st.element.pPr
                if ppr is not None and ppr.numPr is not None:
                    base = ppr.numPr.numId.val
        abs_id = None
        for n in numbering.findall(qn("w:num")):
            if int(n.get(qn("w:numId"))) == base:
                abs_id = n.find(qn("w:abstractNumId")).get(qn("w:val"))
        ids = [int(n.get(qn("w:numId"))) for n in numbering.findall(qn("w:num"))]
        new_id = max(ids) + 1
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(new_id))
        a = OxmlElement("w:abstractNumId")
        a.set(qn("w:val"), abs_id)
        num.append(a)
        ov = OxmlElement("w:lvlOverride")
        ov.set(qn("w:ilvl"), "0")
        so = OxmlElement("w:startOverride")
        so.set(qn("w:val"), "1")
        ov.append(so)
        num.append(ov)
        numbering.append(num)
        return new_id

    def table(self, header, rows, widths=None):
        t = self.doc.add_table(rows=1, cols=len(header))
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for c, h in zip(t.rows[0].cells, header):
            c.text = ""
            r = c.paragraphs[0].add_run(h)
            r.bold = True
            _shade(c, "D9D9D9")
        for row in rows:
            cells = t.add_row().cells
            for c, v in zip(cells, row):
                c.text = ""
                _runs(c.paragraphs[0], v)
        if widths:
            for row in t.rows:
                for c, w in zip(row.cells, widths):
                    c.width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def box(self, label, text, fill="F2F2F2"):
        """Shaded one-cell callout. Lines split on newline."""
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        c = t.rows[0].cells[0]
        _shade(c, fill)
        c.text = ""
        first = c.paragraphs[0]
        r = first.add_run(label)
        r.bold = True
        for line in text.split("\n"):
            par = c.add_paragraph()
            _runs(par, line)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def memory(self, text):
        self.box("Memory Aid", text, "E8F0FE")

    def watch(self, text):
        self.box("Watch Out", text, "FDECEA")

    def pagebreak(self):
        self.doc.add_page_break()

    def save(self, path):
        self.doc.save(path)
        pdf = os.path.splitext(path)[0] + ".pdf"
        ps = (
            "$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
            f"$d=$w.Documents.Open('{path}');$d.SaveAs([ref]'{pdf}',[ref]17);$d.Close();$w.Quit()"
        )
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)
        return pdf
